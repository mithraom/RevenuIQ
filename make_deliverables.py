"""Builds report (.docx), slides (.pptx), viva notes and figures from YOUR data/train.csv results.
Run:  py make_deliverables.py     (needs the real Kaggle train.csv; delete data/.synthetic first)"""
import warnings; warnings.filterwarnings("ignore")
import numpy as np, pandas as pd, matplotlib; matplotlib.use("Agg")
import matplotlib.pyplot as plt
from pathlib import Path
from sklearn.model_selection import cross_val_predict
import pipeline as pl
NAME, ROLL, COURSE = "Mithra Omprakash", "2403717673722027", "22MDC56 Machine Learning Laboratory"
OUT = Path("deliverables"); FIG = OUT / "figures"; FIG.mkdir(parents=True, exist_ok=True)
if (Path("data") / ".synthetic").exists(): raise SystemExit("data/.synthetic exists: put the REAL train.csv in data/ and delete that file first.")
raw = pl.load_raw(); X, y = pl.preprocess(raw); A = pl.train_and_save(); res, best, imp = A["results"], A["best"], A["importance"]; S = pl.summary(raw, X, y)
M = lambda v: f"${v/1e6:.2f}M"; row = lambda n: res[res.Model == n].iloc[0]
b = res.loc[0]; base = row("Baseline (mean)"); gain = (base.RMSE - b.RMSE) / base.RMSE * 100
rf, prf, lr = row("Random Forest"), row("PCA + Random Forest"), row("Linear Regression")
DARK, CL, GD = "#1F1B16", "#B4492D", "#C9A24B"

# ---------- figures ----------
def save(name): plt.tight_layout(); plt.savefig(FIG / name, dpi=170); plt.close()
f, ax = plt.subplots(1, 2, figsize=(9, 3.4)); ax[0].hist(y / 1e6, 25, color=CL); ax[0].set(title="Revenue ($M): right-skewed"); ax[1].hist(np.log1p(y), 25, color=GD); ax[1].set(title="log(revenue): near-normal"); save("fig1_revenue_distribution.png")
plt.figure(figsize=(7, 3.8)); r = res.iloc[::-1]; plt.barh(r.Model, r.RMSE / 1e6, xerr=r.RMSE_sd / 1e6, color=[GD if m == best else "#bbb" for m in r.Model]); plt.xlabel("CV RMSE ($M), lower is better"); plt.title("Model comparison (5x3 repeated CV)"); save("fig2_model_comparison.png")
pr = cross_val_predict(pl.get_model(best), X, y, cv=5); plt.figure(figsize=(4.8, 4.4)); plt.scatter(y / 1e6, pr / 1e6, c=CL, s=18); m = y.max() / 1e6; plt.plot([0, m], [0, m], "--", c=GD); plt.xlabel("Actual ($M)"); plt.ylabel("Predicted ($M)"); plt.title(f"Out-of-fold: {best}"); save("fig3_pred_vs_actual.png")
t = imp.head(10)[::-1]; plt.figure(figsize=(6, 3.8)); plt.barh(t.index, t.values, color=CL); plt.title("Random Forest feature importance (top 10)"); save("fig4_importance.png")
cols = X.assign(revenue=y).corr()["revenue"].abs().sort_values(ascending=False).index[:12]; c = X.assign(revenue=y).corr().loc[cols, cols]
plt.figure(figsize=(6, 5)); plt.imshow(c, cmap="RdBu_r", vmin=-1, vmax=1); plt.xticks(range(12), cols, rotation=90, fontsize=7); plt.yticks(range(12), cols, fontsize=7); plt.colorbar(); plt.title("Correlation heatmap (top 12 by |r| with revenue)"); save("fig5_correlation.png")
f, ax = plt.subplots(1, 2, figsize=(9, 3.4)); raw.assign(r=y / 1e6).boxplot("r", by="City Group", ax=ax[0]); raw.assign(r=y / 1e6).boxplot("r", by="Type", ax=ax[1]); ax[0].set(title="By city group", ylabel="$M"); ax[1].set(title="By type"); plt.suptitle(""); save("fig6_segments.png")
bigm = y[raw["City Group"] == "Big Cities"].median(); othm = y[raw["City Group"] != "Big Cities"].median(); age_r = np.corrcoef(X.days_open, y)[0, 1]; top = [i for i in imp.index if i.startswith("P")][:3]

# ---------- text ----------
verdict = ("clearly beats the mean-only baseline" if gain > 10 else "only marginally beats the mean-only baseline" if gain > 0 else "does not beat the mean-only baseline")
r2txt = (f"R² of {b.R2:.2f} means the model explains a modest share of the variance" if b.R2 > .1 else f"R² of {b.R2:.2f} means the model explains almost none of the variance")
pcatxt = (f"PCA + Random Forest ({M(prf.RMSE)}) was {'better' if prf.RMSE < rf.RMSE else 'worse'} than Random Forest on all features ({M(rf.RMSE)}), so PCA {'helped' if prf.RMSE < rf.RMSE else 'did not help'} on this dataset")
ABSTRACT = (f"RevenuIQ is a site-evaluation decision-support system for restaurant chains. Using the Kaggle TFI dataset ({S['rows']} branches, {S['features']} engineered features) it compares eight models under 5x3 repeated cross-validation. The best model, {best}, reached a cross-validated RMSE of {M(b.RMSE)} and R² of {b.R2:.2f} and {verdict} (RMSE {M(base.RMSE)}). The model is deployed behind a Streamlit interface that converts a forecast into an open/review/avoid verdict supported by comparable branches and lever sensitivity.")
SURVEY = [("[1] Lan et al., 2022", "Combines semantic, temporal and spatial context with supervised models for event attendance; multi-feature models beat single-feature baselines.", "Motivated adding contextual and operational predictors beyond raw numbers."),
 ("[2] Nguyen et al., 2022", "XGBoost and SVM outperform logistic regression on ticket-holder attendance (5,946 holders).", "Justified testing ensembles against linear baselines."),
 ("[3] Ooi & Mafas", "Review of attendance predictors (XGBoost, SVM, contextual models); weather and promotion recur as predictors.", "Guided feature selection thinking."),
 ("[4] ACM, 2020", "LSTM on social-media text predicts attendance intent, 90-92% accuracy.", "Shows value of leading indicators; no equivalent field exists in TFI, so not reproduced."),
 ("[5] US Patents 11,790,386 / 12,288,220", "Machine-learned attendance prediction drives ticket allocation.", "Shows operational value of accurate prediction."),
 ("[6] JCIT", "MLR, Decision Tree, RF and SVR on restaurant branch revenue (1996-2014); SVR lowest error.", "Added SVR to the model set."),
 ("[7] ECAI, 2024", "Correlation and p-value feature selection before MLR and RF on restaurant revenue.", "Matches our correlation-filter step."),
 ("[8] Gogolev & Sotnikov, 2020", "OLS, Elastic Net, SVR, RF on a chain's revenue (5,889 observations); handles correlated features.", "Highlighted correlated-feature handling."),
 ("[9] Rao et al., 2021", "PCA + Random Forest on the same Kaggle TFI dataset.", "Direct basis for our PCA + RF pipeline."),
 ("[10] Rao et al.", "Linear regression and tree ensembles on 40+ correlated predictors; trees cope better.", "Reinforced RF vs linear comparison.")]
GAPS = ["Most studies apply regression, tree ensembles or deep learning individually and rarely combine dimensionality reduction, a full model comparison and a usable interface.",
 "Many papers stop at model accuracy and do not deliver an interactive decision-support tool.",
 f"The TFI dataset ({S['features']} correlated predictors vs {S['rows']} rows) invites overfitting, which few works address through in-fold PCA and correlation filtering.",
 "RevenuIQ addresses these by combining correlation filtering, PCA, an eight-model cross-validated comparison with a baseline, and a Streamlit site-evaluation interface."]
INSIGHTS = [f"Location: median revenue is {M(bigm)} in Big Cities versus {M(othm)} elsewhere.", f"Maturity: revenue correlates {age_r:+.2f} with days open, so budgets should include a ramp-up period.",
 f"Drivers: the most influential anonymised indicators are {', '.join(top)}; they should be prioritised when scouting sites.", f"Uncertainty: the typical error is about {M(b.RMSE)}, so forecasts should rank sites rather than set financial commitments.",
 "Decision rule: sites forecast above the 66th percentile of existing branches are marked OPEN, 33rd-66th REVIEW and below AVOID. These cutoffs are illustrative and should be calibrated with business rules."]
LIMITS = [f"Only {S['rows']} training rows against {S['features']} features limits reliability; cross-validation variance is high (RMSE sd about {M(b.RMSE_sd)}).", "P1-P37 are anonymised, so their business meaning cannot be explained.",
 f"{S['zero_heavy']} P-variables contain more than 30% zeros that may be missing values; they were kept as given.", f"{S['outliers']} outlier branches were kept because they are real observations; log transform reduces their influence.",
 "The City column was dropped because of high cardinality; City Group was kept.", "Verdict cutoffs are heuristic. No external validation set was available."]
FUTURE = ["Add real footfall/attendance data (for example the Kaggle Recruit Restaurant Visitor Forecasting set) as a second module.", "Add social-media and weather signals suggested in the literature.", "Calibrate prediction intervals (conformal prediction) and the OPEN/REVIEW/AVOID thresholds.", "Deploy with a feedback loop to retrain as new branches report revenue."]

# ---------- report ----------
from docx import Document
from docx.shared import Pt, Inches
d = Document(); d.styles["Normal"].font.name = "Calibri"; d.styles["Normal"].font.size = Pt(11)
H = lambda t, l=1: d.add_heading(t, l); P = lambda t: d.add_paragraph(t); B = lambda items: [d.add_paragraph(i, style="List Bullet") for i in items]
def TB(rows, hdr):
    t = d.add_table(rows=1, cols=len(hdr)); t.style = "Light Grid Accent 1"
    for i, h in enumerate(hdr): t.rows[0].cells[i].text = h
    for r in rows:
        c = t.add_row().cells
        for i, v in enumerate(r): c[i].text = str(v)
def FG(n, cap, w=6): d.add_picture(str(FIG / n), width=Inches(w)); P(cap).italic = True
d.add_heading("RevenuIQ: Restaurant Revenue Prediction and Site Evaluation", 0); P(f"{COURSE}\nProject report\n{NAME} [{ROLL}]")
H("Abstract"); P(ABSTRACT)
H("1. Problem statement and objectives"); P("Opening a restaurant in an unsuitable location is a costly decision. Chains need a defensible estimate of what a candidate site will earn and a way to compare sites. The problem is to predict branch revenue from location, maturity and anonymised demographic and real-estate indicators, and to present the result as a decision-support tool.")
B(["Build a leakage-free pipeline for a small, high-dimensional dataset.", "Compare linear, tree-ensemble, kernel and boosted models against a baseline.", "Deliver an interactive interface that turns a forecast into a site decision."])
H("2. Literature survey")
TB([(a, b_, c) for a, b_, c in SURVEY], ["Reference", "Summary", "Relevance to this project"]); H("Gap analysis", 2); B(GAPS)
H("3. Dataset"); P(f"The Kaggle TFI Restaurant Revenue Prediction training set has {S['rows']} branches with Open Date, City, City Group, Type, 37 anonymised indicators (P1-P37) and annual revenue. Revenue ranges from {M(S['min'])} to {M(S['max'])} (median {M(S['median'])}, mean {M(S['mean'])}). There are {S['missing']} missing cells.")
H("4. Methodology"); H("4.1 Preprocessing and feature engineering", 2)
B([f"Open Date converted to days open (reference date 23 March 2015).", "City Group encoded as 0/1; Type one-hot encoded; City dropped (too many distinct values for 137 rows).",
   f"Revenue is right-skewed (skewness {S['skew_raw']:.2f}); log transform reduces it to {S['skew_log']:.2f}. Models predict log revenue and results are converted back.",
   f"{S['zero_heavy']} indicators have over 30% zeros, likely missing values; kept as given and handled by scaling and PCA. {S['outliers']} IQR outliers were retained as genuine branches.",
   f"Correlation filter (|r| > 0.9) reduces {S['features']} features to {S['after_corr']}; PCA keeps 95% variance with {S['pca95']} components. Both are fitted inside each cross-validation fold to prevent leakage."])
H("4.2 Exploratory data analysis", 2); FG("fig1_revenue_distribution.png", "Figure 1: Revenue before and after log transform."); FG("fig6_segments.png", "Figure 2: Revenue by city group and type."); FG("fig5_correlation.png", "Figure 3: Correlation heatmap.", 4.8)
H("4.3 Models and validation", 2); P("Eight models were evaluated with 5-fold cross-validation repeated 3 times (RMSE, MAE, R²): mean baseline, Linear Regression, Linear Regression + PCA, Random Forest, Random Forest (tuned by nested grid search over depth and leaf size), PCA + Random Forest, SVR + PCA and XGBoost. The best-ranked real model is refitted on all data for deployment.")
H("5. System design"); P("pipeline.py holds preprocessing, models and cross-validation. The Streamlit app (app.py and views/) loads the saved model and offers five pages: Home, Site Evaluator (forecast, percentile, verdict, comparable branches, lever sensitivity, maturity curve, downloadable report), Compare & Shortlist (A vs B and batch CSV scoring), Analytics (interactive EDA) and Model Lab.")
H("6. Results and analysis")
TB([(r.Model, f"{r.RMSE/1e6:.3f}", f"{r.MAE/1e6:.3f}", f"{r.R2:.3f}") for r in res.itertuples()], ["Model", "RMSE ($M)", "MAE ($M)", "R²"]); FG("fig2_model_comparison.png", "Figure 4: Cross-validated RMSE by model.", 5.5)
P(f"{best} achieved the lowest cross-validated RMSE ({M(b.RMSE)}), which {verdict} ({M(base.RMSE)}), an improvement of {gain:.1f}%. An {r2txt}. {pcatxt}. Linear Regression reached {M(lr.RMSE)}, so tree-based models {'outperformed' if min(rf.RMSE, b.RMSE) < lr.RMSE else 'did not outperform'} the linear baseline, consistent with the literature on correlated predictors.")
FG("fig3_pred_vs_actual.png", "Figure 5: Out-of-fold predicted vs actual revenue.", 4.2); FG("fig4_importance.png", "Figure 6: Feature importance.", 5)
H("7. Business insights"); B(INSIGHTS); H("8. Limitations", 1); B(LIMITS); H("9. Conclusion and future work"); P(f"RevenuIQ shows how a small, noisy dataset can still support a rigorous, leakage-free modelling workflow and an interface that frames the output as a site decision. The honest finding is that the data supports ranking sites more than precise forecasting. {b.R2:.2f} R² and {M(b.RMSE)} RMSE are reported as measured.")
B(FUTURE); H("References")
for i, r in enumerate(["C. Lan et al., Research on the Prediction System of Event Attendance in an Event-Based Social Network, Wireless Communications and Mobile Computing, 2022.", "Nguyen et al., Predicting Individual Event Attendance with Machine Learning, Applied Economics 54(27), 2022.", "J.L. Ooi and R. Mafas, A Review on the Predictive Algorithms for Event Attendance Prediction in Event-Based Social Networks.", "Event Attendance Prediction Using Social Media, ACM Digital Library, 2020.", "US Patents 11,790,386 / 12,288,220, Machine-Learned Attendance Prediction for Ticket Distribution.", "Predicting Restaurant Revenue using Machine Learning, JCIT.", "Restaurant Revenue Prediction Using ML: Multiple Linear Regression and Random Forest, ECAI 2024.", "S. Gogolev and E. Sotnikov, Comparison of ML Algorithms in Restaurant Revenue Prediction, 2020.", "G.P.S. Rao et al., Machine Learning based Restaurant Revenue Prediction, Springer AISC, 2021.", "G.S. Rao, K.A. Shastry, S. Sathyashree et al., Restaurant Revenue Prediction Applying Supervised Learning Methods."], 1): P(f"[{i}] {r}")
P("Dataset: Kaggle, TFI Restaurant Revenue Prediction competition. GitHub: <add your repository link here>")
d.save(OUT / "RevenuIQ_Report.docx")

# ---------- slides ----------
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches as I, Pt as Pp
pr_ = Presentation(); pr_.slide_width, pr_.slide_height = I(13.333), I(7.5)
def slide(title, bullets=(), pic=None, pw=5.6, small=False):
    s = pr_.slides.add_slide(pr_.slide_layouts[6]); s.background.fill.solid(); s.background.fill.fore_color.rgb = RGBColor(14, 11, 9)
    tb = s.shapes.add_textbox(I(.6), I(.4), I(12), I(1)).text_frame; tb.text = title; p = tb.paragraphs[0]; p.font.size = Pp(34); p.font.bold = True; p.font.color.rgb = RGBColor(217, 180, 90)
    w = 6.4 if pic else 12; bx = s.shapes.add_textbox(I(.6), I(1.5), I(w), I(5.5)).text_frame; bx.word_wrap = True
    for i, t in enumerate(bullets):
        p = bx.paragraphs[0] if i == 0 else bx.add_paragraph(); p.text = "• " + t; p.font.size = Pp(15 if small else 19); p.font.color.rgb = RGBColor(243, 235, 221); p.space_after = Pp(8)
    if pic: s.shapes.add_picture(str(FIG / pic), I(7.2), I(1.6), width=I(pw))
slide("RevenuIQ", ["Restaurant revenue prediction and site evaluation", f"{NAME} [{ROLL}]", COURSE])
slide("Problem and objectives", ["Wrong site choice is expensive; chains need defensible revenue forecasts", "Predict branch revenue from location, maturity and anonymised indicators", "Deliver a decision-support tool, not just a model"])
slide("Literature survey and gaps", GAPS, small=True)
slide("Dataset (Kaggle TFI)", [f"{S['rows']} branches, {S['features']} engineered features", f"Revenue {M(S['min'])} to {M(S['max'])}, median {M(S['median'])}", "P1-P37 anonymised; City Group, Type, Open Date"], pic="fig1_revenue_distribution.png", pw=5.8)
slide("Preprocessing", [f"Days open from Open Date; log revenue (skew {S['skew_raw']:.1f} to {S['skew_log']:.1f})", f"{S['zero_heavy']} zero-heavy indicators kept; {S['outliers']} outliers retained", f"Correlation filter: {S['features']} to {S['after_corr']}; PCA 95% = {S['pca95']} components", "All fitted inside CV folds: no leakage"], small=True)
slide("EDA", ["Big Cities earn more than other cities", f"Revenue vs maturity correlation {age_r:+.2f}", "Several indicators are highly correlated"], pic="fig6_segments.png", pw=5.8)
slide("Models and validation", ["Baseline, Linear, Linear+PCA, Random Forest, tuned RF, PCA+RF, SVR+PCA, XGBoost", "5-fold CV repeated 3 times; RMSE, MAE, R²", "Nested grid search for tuned RF"])
slide("Results", [f"Best: {best}", f"RMSE {M(b.RMSE)}, R² {b.R2:.2f}", f"Baseline RMSE {M(base.RMSE)}; improvement {gain:.1f}%", pcatxt], pic="fig2_model_comparison.png", pw=5.8, small=True)
slide("Predicted vs actual and drivers", [f"Top indicators: {', '.join(top)}", "Points near the dashed line are accurate", "Model ranks sites better than it prices them"], pic="fig3_pred_vs_actual.png", pw=4.6)
slide("System and UI", ["Streamlit multi-page app over a saved model", "Site Evaluator: forecast, percentile, OPEN/REVIEW/AVOID, peers, levers", "Compare & Shortlist: A vs B and batch CSV", "Analytics and Model Lab pages"])
slide("Business insights", INSIGHTS, small=True)
slide("Limitations and future work", LIMITS[:4] + FUTURE[:2], small=True)
slide("Conclusion", ["Leakage-free pipeline on a small noisy dataset", f"{best} is the best model, honestly reported (R² {b.R2:.2f})", "Interface turns forecasts into site decisions", "Thank you. Questions?"])
pr_.save(OUT / "RevenuIQ_Slides.pptx")

# ---------- viva ----------
QA = [("What problem does the project solve?", "Choosing restaurant sites. It forecasts branch revenue and turns it into an open/review/avoid verdict."),
 ("Why log-transform revenue?", f"Revenue is right-skewed (skew {S['skew_raw']:.2f}); the log brings it to {S['skew_log']:.2f} so large branches do not dominate."),
 ("Why only 137 rows and is that a problem?", "That is the dataset size. It causes high variance, which is why I used repeated cross-validation and report standard deviations."),
 ("How did you avoid data leakage?", "Scaling, correlation filtering and PCA live inside the sklearn Pipeline, so they are fitted on training folds only."),
 ("Why PCA?", f"{S['features']} correlated predictors on {S['rows']} rows risk overfitting; PCA keeps 95% variance with {S['pca95']} components. {pcatxt}."),
 ("Did PCA help?", pcatxt + ". I report it as measured, not as hoped."),
 ("Which model won and why?", f"{best}: RMSE {M(b.RMSE)}, R² {b.R2:.2f}. Tree ensembles handle correlated features and non-linear effects better than linear models."),
 ("How do you know the model is useful?", f"It is compared with a mean-only baseline (RMSE {M(base.RMSE)}); improvement is {gain:.1f}%."),
 ("Why is R² low?", "Small sample, anonymised noisy indicators and outliers. The honest conclusion is that the data supports ranking better than exact forecasting."),
 ("Why not use City?", "Too many distinct cities for 137 rows; City Group captures the main location effect."),
 ("What about zeros in P-variables?", f"{S['zero_heavy']} have over 30% zeros, probably missing. I kept them rather than inventing imputations; scaling and PCA absorb them. Imputation is future work."),
 ("What about outliers?", f"{S['outliers']} outliers are real high-revenue branches; removing them would hide real behaviour. Log transform limits their effect."),
 ("What is RMSE vs MAE vs R²?", "RMSE penalises large errors, MAE is the average absolute error, R² is the share of variance explained versus predicting the mean."),
 ("Why repeated CV?", "With 137 rows a single split is unreliable; 5x3 repeats average over 15 splits."),
 ("How was tuning done?", "Grid search over max depth and minimum leaf size inside each outer fold (nested), so tuning does not leak."),
 ("What does the OPEN/REVIEW/AVOID verdict mean?", "It uses percentile of the forecast among existing branches (66th and 33rd). The cutoffs are illustrative, not validated business rules."),
 ("What are comparable branches?", "The five nearest real branches in the space of top indicators, days open and city group; they sanity-check the forecast."),
 ("What does lever sensitivity show?", "How the forecast changes when one key indicator moves from its 10th to 90th percentile."),
 ("Why is there no attendance prediction?", "TFI contains no attendance data. The survey papers motivated contextual features; real footfall data (for example Recruit Restaurant Visitor Forecasting) is future work."),
 ("What are the main limitations?", " ".join(LIMITS[:3])),
 ("What would you improve?", " ".join(FUTURE[:3])),
 ("How is the model deployed in the app?", "Training happens once and is saved with joblib; the Streamlit pages load it and call predict on the user's inputs.")]
(OUT / "viva_notes.md").write_text("# Viva preparation\n\n" + "\n\n".join(f"**Q{i}. {q}**\n{a}" for i, (q, a) in enumerate(QA, 1)) + "\n")
print("Done. See deliverables/: RevenuIQ_Report.docx, RevenuIQ_Slides.pptx, viva_notes.md, figures/")
